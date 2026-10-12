import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path
from time import monotonic

import pytest
from docx import Document
from pypdf import PdfWriter

import laplace.tools.documents  # noqa: F401
from laplace.config import Settings
from laplace.llm.base import BudgetStop
from laplace.services.budget import TurnBudget
from laplace.services.documents import DocumentLimits, extract_document
from laplace.tools.base import ToolContext, ToolFailure, execute


def _limits() -> DocumentLimits:
    return DocumentLimits(
        max_bytes=2_000_000,
        max_units=50,
        max_unit_chars=1_000,
        max_page_chars=1_000,
        max_expanded_bytes=5_000_000,
        parse_timeout_seconds=5,
        memory_limit_bytes=256 * 1024 * 1024,
    )


def _budget() -> TurnBudget:
    return TurnBudget(settings=Settings(_env_file=None))


def _context(path, attachment_id="opaque-1", *, budget=None) -> ToolContext:
    return ToolContext(
        user_id=1,
        telegram_user_id=10,
        transport="telegram",
        is_private_chat=True,
        conversation_id=1,
        request_message_id=1,
        task_id=1,
        step_index=1,
        now_utc=datetime(2026, 1, 1, tzinfo=UTC),
        timezone="UTC",
        budget=budget if budget is not None else _budget(),
        attachments={attachment_id: path},
    )


def test_txt_units_and_tool_sources_do_not_disclose_path(tmp_path, monkeypatch) -> None:
    path = tmp_path / "private-name.txt"
    path.write_text("alpha\nbeta\ngamma", encoding="utf-8")
    document = extract_document({"opaque-1": path}, "opaque-1", _limits(), budget=_budget())
    assert [(unit.locator, unit.text) for unit in document.units] == [
        ("line:1", "alpha"),
        ("line:2", "beta"),
        ("line:3", "gamma"),
    ]

    monkeypatch.setenv("LAPLACE_DOCUMENT_FULL_SUMMARY_MAX_CHARS", "1000")
    result = execute(
        "read_document", {"attachment_id": "opaque-1", "start_unit": 1}, _context(path)
    )
    assert result.ok is True
    assert result.sources == [
        "attachment:opaque-1#line:1",
        "attachment:opaque-1#line:2",
        "attachment:opaque-1#line:3",
    ]
    assert str(tmp_path) not in result.model_dump_json()


def test_read_document_paginates_whole_locator_units(tmp_path, monkeypatch) -> None:
    path = tmp_path / "long.txt"
    path.write_text("\n".join([("a" * 600), ("b" * 600), ("c" * 600)]))
    monkeypatch.setenv("LAPLACE_DOCUMENT_FULL_SUMMARY_MAX_CHARS", "1000")

    first = execute("read_document", {"attachment_id": "opaque-1", "start_unit": 1}, _context(path))
    second = execute(
        "read_document", {"attachment_id": "opaque-1", "start_unit": 2}, _context(path)
    )

    assert [item["locator"] for item in first.data["units"]] == ["line:1"]
    assert (first.data["next_unit"], first.data["complete"]) == (2, False)
    assert [item["locator"] for item in second.data["units"]] == ["line:2"]
    assert second.data["next_unit"] == 3


def test_txt_rejects_excess_units_during_iteration(tmp_path) -> None:
    path = tmp_path / "many-lines.txt"
    path.write_text("x\n" * 10_000, encoding="utf-8")

    with pytest.raises(ToolFailure) as raised:
        extract_document({"txt": path}, "txt", _limits(), budget=_budget())
    assert raised.value.error_code == "unsupported_content"


def test_docx_paragraph_and_table_locators(tmp_path) -> None:
    path = tmp_path / "sample.docx"
    source = Document()
    source.add_paragraph("Paragraph text")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    source.save(path)

    document = extract_document({"doc": path}, "doc", _limits(), budget=_budget())
    assert [unit.locator for unit in document.units] == ["paragraph:1", "table:1/row:1"]
    assert document.units[1].text == "A | B"


def test_docx_expanded_size_limit_is_checked_before_parsing(tmp_path) -> None:
    path = tmp_path / "large.docx"
    source = Document()
    source.add_paragraph("content")
    source.save(path)
    limits = DocumentLimits(
        max_bytes=2_000_000,
        max_units=50,
        max_unit_chars=1_000,
        max_page_chars=1_000,
        max_expanded_bytes=10,
    )

    with pytest.raises(ToolFailure) as raised:
        extract_document({"doc": path}, "doc", limits, budget=_budget())
    assert raised.value.error_code == "unsupported_content"


def test_image_only_pdf_is_explicitly_unsupported(tmp_path) -> None:
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as stream:
        writer.write(stream)

    with pytest.raises(ToolFailure) as raised:
        extract_document({"pdf": path}, "pdf", _limits(), budget=_budget())
    assert raised.value.error_code == "unsupported_content"


@pytest.mark.parametrize("filename", ["bad.pdf", "bad.docx", "binary.txt"])
def test_corrupt_or_binary_inputs_fail_cleanly(tmp_path, filename: str) -> None:
    path = tmp_path / filename
    path.write_bytes(b"not a document\x00")
    with pytest.raises(ToolFailure) as raised:
        extract_document({"id": path}, "id", _limits(), budget=_budget())
    assert raised.value.error_code == "unsupported_content"


def test_unknown_and_stale_attachment_are_indistinguishable(tmp_path) -> None:
    missing = tmp_path / "gone.txt"
    messages = []
    for paths in ({}, {"id": missing}):
        with pytest.raises(ToolFailure) as raised:
            extract_document(paths, "id", _limits(), budget=_budget())
        messages.append((raised.value.error_code, raised.value.message))
    assert messages[0] == messages[1] == ("not_found", "Attachment was not found.")


@pytest.mark.parametrize("limiter", ["parser", "turn"])
def test_parser_timeout_kills_and_reaps_actual_worker(tmp_path, monkeypatch, limiter) -> None:
    from dataclasses import replace

    import laplace.services.documents as documents

    path = tmp_path / "slow.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    processes = []
    original_popen = subprocess.Popen

    def slow_worker(_command, **kwargs):
        process = original_popen(
            [
                sys.executable,
                "-c",
                "import sys,time; sys.stdin.buffer.read(); time.sleep(30)",
            ],
            **kwargs,
        )
        processes.append(process)
        return process

    monkeypatch.setattr(documents.subprocess, "Popen", slow_worker)
    offset = [0.0]
    budget = TurnBudget(
        settings=Settings(_env_file=None, max_seconds=1),
        clock=lambda: monotonic() + offset[0],
    )
    limits = _limits()
    if limiter == "parser":
        limits = replace(limits, parse_timeout_seconds=0.08)
    else:
        offset[0] = 0.92
    started = monotonic()
    if limiter == "turn":
        with pytest.raises(BudgetStop) as stopped:
            extract_document({"pdf": path}, "pdf", limits, budget=budget)
        assert stopped.value.reason == "time_limit"
    else:
        with pytest.raises(ToolFailure) as failed:
            extract_document({"pdf": path}, "pdf", limits, budget=budget)
        assert (failed.value.error_code, failed.value.retryable) == ("timeout", True)
    assert monotonic() - started < 0.8
    assert len(processes) == 1
    assert processes[0].poll() is not None and processes[0].returncode < 0


def test_document_read_finishing_at_deadline_is_not_returned(tmp_path, monkeypatch) -> None:
    path = tmp_path / "late.txt"
    path.write_text("late evidence", encoding="utf-8")
    elapsed = [0.0]
    budget = TurnBudget(
        settings=Settings(_env_file=None, max_seconds=1), clock=lambda: elapsed[0]
    )
    original_read = Path.read_bytes

    def read_at_deadline(self):
        payload = original_read(self)
        elapsed[0] = 1.0
        return payload

    monkeypatch.setattr(Path, "read_bytes", read_at_deadline)
    with pytest.raises(BudgetStop) as stopped:
        execute("read_document", {"attachment_id": "opaque-1"}, _context(path, budget=budget))
    assert stopped.value.reason == "time_limit"


def test_expired_document_budget_prevents_attachment_access(tmp_path, monkeypatch) -> None:
    path = tmp_path / "secret.txt"
    elapsed = [0.0]
    budget = TurnBudget(
        settings=Settings(_env_file=None, max_seconds=1), clock=lambda: elapsed[0]
    )
    accessed = []
    monkeypatch.setattr(Path, "stat", lambda self: accessed.append(self))
    elapsed[0] = 1.0
    with pytest.raises(BudgetStop) as stopped:
        extract_document({"id": path}, "id", _limits(), budget=budget)
    assert stopped.value.reason == "time_limit"
    assert accessed == []
